# -*- coding: utf-8 -*-
"""PM 待办助手 使用手册 —— 图文并茂 Word 文档生成脚本"""
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
import numpy as np
import os

# ── 中文字体 ──
plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei"]
plt.rcParams["axes.unicode_minus"] = False

WORK = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(WORK, "assets", "manual_imgs")
os.makedirs(IMG, exist_ok=True)

# ── 配色 ──
C_DEEP   = "#1B3A5B"   # 深蓝
C_TEAL   = "#2A9D8F"   # 青绿
C_ORANGE = "#E76F51"   # 橙红
C_YELLOW = "#E9C46A"   # 暖黄
C_LIGHT  = "#F4F6F8"   # 浅灰
C_GREY   = "#8A99A8"
C_PURPLE = "#7B68EE"

# ============================================================
#  图1 — 封面横幅
# ============================================================
def make_banner():
    fig, ax = plt.subplots(figsize=(7, 2.4))
    ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
    # 渐变背景
    grad = np.linspace(0, 1, 256).reshape(1, -1)
    ax.imshow(grad, extent=[0,1,0,1], aspect="auto", cmap="GnBu", zorder=0)
    ax.text(0.5, 0.62, "PM 待办助手", ha="center", va="center",
            fontsize=28, fontweight="bold", color="white")
    ax.text(0.5, 0.30, "本地优先的项目管理 / 待办工具  ·  使用手册",
            ha="center", va="center", fontsize=12, color="#E0F0E8")
    fig.tight_layout(pad=0)
    p = os.path.join(IMG, "banner.png"); fig.savefig(p, dpi=200); plt.close(fig)
    return p

# ============================================================
#  图2 — 侧栏四分组结构图
# ============================================================
def make_sidebar():
    fig, ax = plt.subplots(figsize=(5.2, 5.6))
    ax.set_xlim(0, 10); ax.set_ylim(0, 11); ax.axis("off")
    # 外框
    ax.add_patch(FancyBboxPatch((0.3, 0.3), 9.4, 10.4,
                 boxstyle="round,pad=0.15", facecolor=C_LIGHT,
                 edgecolor=C_GREY, linewidth=1.2))
    # Logo
    ax.text(5, 10.2, "PM 待办助手 v2.0", ha="center", va="center",
            fontsize=12, fontweight="bold", color=C_DEEP)
    ax.plot([0.8, 9.2], [9.8, 9.8], color=C_GREY, lw=0.8)
    # 分组1
    _group(ax, 1.0, 8.4, "今日聚焦 / 收件箱", C_TEAL,
           ["[日] 今日聚焦         8", "[入] 收件箱           3"])
    # 分组2
    _group(ax, 1.0, 6.2, "项目  [+]", C_DEEP,
           ["· 闸机二期       12", "· 示例项目        4", "> 已归档 (2)"])
    # 分组3
    _group(ax, 1.0, 3.8, "洞察", C_PURPLE,
           ["[目] 智能视图（含待清理）", "[灯] 需求池"])
    # 分组4
    _group(ax, 1.0, 1.8, "记录", C_ORANGE,
           ["[会] 会议纪要", "[人] 干系人", "[笔] 每日笔记"])
    # 页脚
    ax.text(5, 0.7, "[主题] [提醒] [回收站] [设置]    仅 4 个高频按钮",
            ha="center", va="center", fontsize=7.5, color=C_GREY)
    fig.tight_layout(pad=0.3)
    p = os.path.join(IMG, "sidebar.png"); fig.savefig(p, dpi=200); plt.close(fig)
    return p

def _group(ax, x, y, title, color, items):
    ax.text(x, y+0.35, title, fontsize=9, fontweight="bold", color=color)
    for i, item in enumerate(items):
        ax.text(x+0.3, y-i*0.38, item, fontsize=8, color="#333333")
    ax.plot([x-0.2, x+8.5], [y-0.15, y-0.15], color="#DDE3EA", lw=0.6)

# ============================================================
#  图3 — 任务四种视图示意
# ============================================================
def make_views():
    fig, axes = plt.subplots(2, 2, figsize=(7.2, 5.0))
    fig.suptitle("任务页四种视图", fontsize=14, fontweight="bold", color=C_DEEP, y=0.98)
    # 列表
    ax = axes[0,0]; ax.set_title("列表视图", fontsize=10, color=C_DEEP)
    ax.set_xlim(0,10); ax.set_ylim(0,6); ax.axis("off")
    for i, (txt, c) in enumerate([("P0 闸机联调  @李四","#E76F51"),("P1 接口对接  @王五","#E9C46A"),
                                   ("P2 文档评审  @张三","#2A9D8F"),("P0 现场部署  @李四","#E76F51")]):
        y = 5 - i*1.3
        ax.add_patch(FancyBboxPatch((0.5,y-0.4),9,0.9, boxstyle="round,pad=0.1",
                     facecolor="white", edgecolor=c, linewidth=1.5))
        ax.text(1, y, txt, fontsize=8, va="center", color="#333")
    # 看板
    ax = axes[0,1]; ax.set_title("看板视图", fontsize=10, color=C_DEEP)
    ax.set_xlim(0,12); ax.set_ylim(0,6); ax.axis("off")
    cols = [("待办", C_GREY, 2), ("进行中", C_TEAL, 1), ("已完成", "#4CAF50", 2)]
    for ci, (name, c, cnt) in enumerate(cols):
        cx = 0.5 + ci*4
        ax.add_patch(FancyBboxPatch((cx,4.8),3.2,0.8, boxstyle="round,pad=0.1",
                     facecolor=c, edgecolor="none"))
        ax.text(cx+1.6, 5.2, f"{name} ({cnt})", ha="center", va="center",
                fontsize=8, color="white", fontweight="bold")
        for k in range(cnt):
            ax.add_patch(FancyBboxPatch((cx+0.15, 4-k*1.3-0.8), 2.9, 1.0,
                         boxstyle="round,pad=0.08", facecolor="white",
                         edgecolor="#CCC", linewidth=0.8))
            ax.text(cx+0.4, 4-k*1.3-0.3, f"任务{ci}{k+1}", fontsize=7, va="center")
    # 日历
    ax = axes[1,0]; ax.set_title("日历视图", fontsize=10, color=C_DEEP)
    days = list(range(1,31))
    cal = np.zeros((5,7))
    d = 1
    for r in range(5):
        for c in range(7):
            if d <= 30: cal[r,c] = d; d += 1
    ax.imshow(cal, cmap="Blues", aspect="auto")
    for r in range(5):
        for c in range(7):
            if cal[r,c] > 0:
                ax.text(c, r, str(int(cal[r,c])), ha="center", va="center",
                        fontsize=7, color="#333" if cal[r,c] < 20 else "white")
    # 标红几个截止日
    for (r,c) in [(1,2),(2,5),(3,1)]:
        ax.add_patch(plt.Rectangle((c-0.5,r-0.5),1,1, fill=False,
                     edgecolor=C_ORANGE, linewidth=2))
    ax.set_xticks([]); ax.set_yticks([])
    # 时间线甘特
    ax = axes[1,1]; ax.set_title("时间线（甘特）", fontsize=10, color=C_DEEP)
    tasks_g = [("需求分析",0,3,C_TEAL),("开发",2,6,C_DEEP),("联调",5,8,C_ORANGE),("上线",7,9,"#4CAF50")]
    for i,(name,s,e,c) in enumerate(tasks_g):
        ax.barh(i, e-s, left=s, height=0.5, color=c, edgecolor="white")
        ax.text(s+0.1, i, name, fontsize=7, va="center", color="white")
    ax.axvline(6, color=C_ORANGE, ls="--", lw=1.2, label="今日")
    ax.set_xlim(0,9); ax.set_ylim(-0.6,3.6)
    ax.set_yticks([]); ax.set_xticks(range(0,10,2))
    ax.tick_params(labelsize=7)
    ax.legend(fontsize=7, loc="upper right")
    fig.tight_layout(pad=1.0, rect=[0,0,1,0.95])
    p = os.path.join(IMG, "views.png"); fig.savefig(p, dpi=200); plt.close(fig)
    return p

# ============================================================
#  图4 — 功能模块分类统计（柱状图）
# ============================================================
def make_feature_chart():
    fig, ax = plt.subplots(figsize=(7, 3.2))
    cats = ["任务与项目", "汇报与沟通", "洞察与统计", "自动化与集成", "组织记忆", "数据安全"]
    counts = [12, 3, 5, 5, 5, 3]
    colors = [C_DEEP, C_TEAL, C_PURPLE, C_ORANGE, C_YELLOW, "#4CAF50"]
    bars = ax.barh(cats, counts, color=colors, edgecolor="white", height=0.6)
    for bar, cnt in zip(bars, counts):
        ax.text(bar.get_width()+0.2, bar.get_y()+bar.get_height()/2,
                str(cnt), va="center", fontsize=9, fontweight="bold", color=C_DEEP)
    ax.set_xlabel("功能数量", fontsize=9)
    ax.set_title("PM 待办助手 — 功能模块分布", fontsize=12, fontweight="bold", color=C_DEEP)
    ax.spines["top"].set_visible(False); ax.spines["right"].set_visible(False)
    ax.tick_params(labelsize=9)
    ax.invert_yaxis()
    fig.tight_layout(pad=0.8)
    p = os.path.join(IMG, "features.png"); fig.savefig(p, dpi=200); plt.close(fig)
    return p

# ============================================================
#  图5 — 典型工作流程图
# ============================================================
def make_workflow():
    fig, ax = plt.subplots(figsize=(7.2, 3.0))
    ax.set_xlim(0, 14); ax.set_ylim(0, 5); ax.axis("off")
    steps = [
        (1.2, 2.5, "快速添加\n/ 收件箱", C_TEAL),
        (4.0, 2.5, "分拣到项目\n设优先级", C_DEEP),
        (6.8, 2.5, "执行计时\n看板/列表", C_ORANGE),
        (9.6, 2.5, "统计预测\n风险告警", C_PURPLE),
        (12.4, 2.5, "日报/周报\n一键发群", "#4CAF50"),
    ]
    for i,(x,y,txt,c) in enumerate(steps):
        ax.add_patch(FancyBboxPatch((x-1.0,y-0.75),2.0,1.5,
                     boxstyle="round,pad=0.12", facecolor=c, edgecolor="none"))
        ax.text(x, y, txt, ha="center", va="center", fontsize=8.5,
                color="white", fontweight="bold")
        if i < len(steps)-1:
            ax.annotate("", xy=(steps[i+1][0]-1.1, y), xytext=(x+1.1, y),
                        arrowprops=dict(arrowstyle="->", color=C_GREY, lw=1.8))
    ax.text(7, 4.6, "典型工作流程", ha="center", fontsize=13,
            fontweight="bold", color=C_DEEP)
    ax.text(7, 0.5, "全流程数据存储在本机 SQLite  ·  AI 仅产草稿，不直接改数据",
            ha="center", fontsize=8, color=C_GREY)
    fig.tight_layout(pad=0.3)
    p = os.path.join(IMG, "workflow.png"); fig.savefig(p, dpi=200); plt.close(fig)
    return p

# ============================================================
#  图6 — 快速添加语法卡片
# ============================================================
def make_quickadd():
    fig, ax = plt.subplots(figsize=(7, 2.4))
    ax.set_xlim(0, 10); ax.set_ylim(0, 5); ax.axis("off")
    ax.add_patch(FancyBboxPatch((0.3,0.3),9.4,4.4,
                 boxstyle="round,pad=0.2", facecolor="#1E2A38", edgecolor=C_TEAL, linewidth=1.5))
    ax.text(5, 4.2, "快速添加语法示例", ha="center", fontsize=11,
            fontweight="bold", color="white")
    ax.text(5, 3.2, "下周三 闸机联调 @李四 !P0 #风险 #每周",
            ha="center", fontsize=15, color=C_YELLOW)
    # 注解
    notes = [
        ("日期", "今天/明天/下周三/+3天/2026-09-14", C_TEAL),
        ("@", "负责人  @张三", C_ORANGE),
        ("!", "优先级  !P0 !P1 !P2", "#E9C46A"),
        ("#", "标记  #风险 #每日 #每周 #每月", C_PURPLE),
    ]
    for i,(k,v,c) in enumerate(notes):
        y = 2.3 - i*0.5
        ax.text(1.5, y, k, fontsize=8, fontweight="bold", color=c)
        ax.text(2.5, y, v, fontsize=8, color="#CCCCCC")
    fig.tight_layout(pad=0.3)
    p = os.path.join(IMG, "quickadd.png"); fig.savefig(p, dpi=200); plt.close(fig)
    return p

# ============================================================
#  生成所有图片
# ============================================================
print("生成配图...")
imgs = {
    "banner":    make_banner(),
    "sidebar":   make_sidebar(),
    "views":     make_views(),
    "features":  make_feature_chart(),
    "workflow":  make_workflow(),
    "quickadd":  make_quickadd(),
}
for k,v in imgs.items():
    print(f"  {k}: {v}")

# ============================================================
#  组装 Word 文档
# ============================================================
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── 页面设置 ──
for sec in doc.sections:
    sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)

# ── 默认样式 ──
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(10.5)
style.font.color.rgb = RGBColor(0x33,0x33,0x33)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

def _set_run(run, name="微软雅黑", size=None, bold=None, color=None, italic=None):
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts"); rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))

def add_h1(text):
    p = doc.add_paragraph()
    p.space_before = Pt(18)
    r = p.add_run(text)
    _set_run(r, size=18, bold=True, color="1B3A5B")
    _hborder(p, "2A9D8F", 1)
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(8)
    return p

def add_h2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_run(r, size=14, bold=True, color="2A9D8F")
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_h3(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_run(r, size=12, bold=True, color="1B3A5B")
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_body(text, size=10.5, color="333333", bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_run(r, size=size, color=color, bold=bold)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = Pt(18)
    return p

def add_bullet(text, color="333333"):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    _set_run(r, size=10, color=color)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = Pt(16)
    return p

def add_image(path, width=6.0, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(path, width=Inches(width))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cp.add_run(caption)
        _set_run(cr, size=9, color="8A99A8", italic=True)
        cp.paragraph_format.space_after = Pt(8)

def add_table(headers, rows, col_widths=None):
    """添加带样式的表格"""
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # 表头
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        _shade(cell, "1B3A5B")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        _set_run(r, size=9.5, bold=True, color="FFFFFF")
    # 数据行
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri+1].cells[ci]
            if ri % 2 == 1:
                _shade(cell, "F4F6F8")
            p = cell.paragraphs[0]
            r = p.add_run(str(val))
            _set_run(r, size=9, color="333333")
    # 列宽
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    # 间距
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hexcolor.lstrip("#"))
    tcPr.append(shd)

def _hborder(p, hexcolor, size=1):
    pPr = p._element.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size*4))
    bottom.set(qn("w:color"), hexcolor.lstrip("#"))
    pbdr.append(bottom)
    pPr.append(pbdr)

def add_callout(text, color="2A9D8F", icon="💡"):
    """彩色提示框"""
    t = doc.add_table(rows=1, cols=1)
    cell = t.rows[0].cells[0]
    _shade(cell, "E8F5F3")
    p = cell.paragraphs[0]
    r = p.add_run(f"{icon}  {text}")
    _set_run(r, size=9.5, color="1B3A5B")
    # 左边框着色
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), color.lstrip("#"))
    borders.append(left)
    tcPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ============================================================
#  封面
# ============================================================
for _ in range(3):
    doc.add_paragraph()
add_image(imgs["banner"], width=6.0)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("v2.1  ·  Tauri 2 + SQLite  ·  100% 本地存储")
_set_run(r, size=11, color="8A99A8")
for _ in range(2):
    doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("面向需要同时盯多个项目、多项任务的项目经理")
_set_run(r, size=12, color="555555")
doc.add_page_break()

# ============================================================
#  目录
# ============================================================
add_h1("目录")
toc_items = [
    "一、产品简介",
    "二、下载安装",
    "三、界面导览",
    "四、核心功能",
    "五、项目管理",
    "六、汇报与沟通",
    "七、洞察与统计",
    "八、自动化与 AI 集成",
    "九、组织记忆",
    "十、数据安全",
    "十一、快捷键速查",
    "十二、从源码构建",
]
for item in toc_items:
    p = doc.add_paragraph()
    r = p.add_run(item)
    _set_run(r, size=11, color="333333")
    p.paragraph_format.space_after = Pt(6)
doc.add_page_break()

# ============================================================
#  一、产品简介
# ============================================================
add_h1("一、产品简介")
add_body("PM 待办助手是一款本地优先的项目管理 / 待办工具，面向需要同时盯多个项目、多项任务的项目经理。")
add_body("基于 Tauri 2 + SQLite 构建，前端为原生 HTML/JS，无任何外部依赖服务，数据 100% 存储在本机。")
add_image(imgs["features"], width=5.8, caption="图：功能模块分布")

add_h2("核心特点")
add_bullet("本地优先：数据存本机 SQLite，启动自动备份（保留最近 30 份），无需联网")
add_bullet("多项目驾驶舱：今日聚焦跨项目汇总，三只青蛙（MIT）固定要事")
add_bullet("四种任务视图：列表 / 看板 / 日历 / 时间线甘特，一键切换")
add_bullet("快速添加语法：自然语言一行搞定日期、负责人、优先级、标记")
add_bullet("AI 提效三件套：粘贴转任务、生成检查清单、日报润色（可选，仅产草稿）")
add_bullet("一键发群：日报 / 周报推送到企业微信 / 钉钉 / 飞书群")
add_bullet("组织记忆：决策日志、会议纪要、干系人跟进、每日笔记")
add_bullet("回收站软删除：删除的数据保留 30 天可恢复")

add_image(imgs["workflow"], width=6.0, caption="图：典型工作流程")

# ============================================================
#  二、下载安装
# ============================================================
add_h1("二、下载安装")
add_h2("Windows 安装")
add_body("1. 前往 GitHub Releases 页面下载安装包：")
add_body("   PMTodoAssistant_x.x.x_x64-setup.exe", color="2A9D8F", bold=True)
add_body("2. 双击安装即可（Windows 10/11，安装时自动处理 WebView2 运行时）")
add_callout("下载地址：github.com/wangjie-git/pm_tool/releases/latest", color="2A9D8F", icon="🔗")

add_h2("系统要求")
add_table(
    ["项目", "要求"],
    [["操作系统", "Windows 10 / 11（64 位）"],
     ["运行时", "WebView2（Win10/11 通常已内置，安装包自动处理）"],
     ["磁盘空间", "约 50 MB（含数据）"],
     ["网络", "无需联网（AI 功能需配置供应商后可选联网）"]],
    col_widths=[2.0, 4.5]
)

# ============================================================
#  三、界面导览
# ============================================================
add_h1("三、界面导览")
add_body("v2.0 对信息架构做了整体重组：全部功能保留，只重排入口，数据模型与后端零改动。")

add_h2("侧栏四分组")
add_image(imgs["sidebar"], width=4.0, caption="图：侧栏结构示意")
add_table(
    ["分组", "功能入口", "说明"],
    [["今日聚焦 / 收件箱", "📅 今日聚焦、📥 收件箱", "跨项目驾驶舱与分拣入口"],
     ["项目", "项目列表 + ＋新建、▸ 已归档", "含归档折叠，点击进入项目工作区"],
     ["洞察", "🎯 智能视图、💡 需求池", "跨项目筛选与需求打分"],
     ["记录", "🗂 会议纪要、👥 干系人、📝 每日笔记", "组织记忆三类入口"]],
    col_widths=[1.8, 2.5, 2.5]
)
add_callout("页脚仅保留 4 个高频按钮：🌙 主题 / ⏰ 提醒 / 🗑 回收站 / ⚙ 设置", color="2A9D8F")

add_h2("项目工作区：两层页签")
add_body("一级页签「任务 / 统计 / 决策 / 档案」是项目下的四个页面；")
add_body("任务页内再用二级分段控件「列表 / 看板 / 日历 / 时间线」切换任务四种视图。")
add_body("顶栏「⋯ 项目工具」菜单统一收拢：Excel 导入向导、导出 .xlsx / CSV / Markdown、项目设置、另存为模板。")

add_h2("设置弹窗：五节导航")
add_table(
    ["节", "内容"],
    [["🧭 通用", "每日收尾问答时间、每日提醒说明、数据安全说明"],
     ["🤝 自动化", "4 个固定规则开关（逾期置顶 / 待清理 / 完成停表 / 收件箱提醒）"],
     ["🖥 桌面集成", "开机自启、全局快速捕获热键、今日目标数、深链接说明"],
     ["🤖 AI 供应商", "供应商卡片 + 预设 / 模型 / 地址 / Key / 拉取模型 / 测试连接"],
     ["💾 数据", "立即备份、导出 JSON、导入 JSON、打开回收站"]],
    col_widths=[1.5, 5.0]
)

# ============================================================
#  四、核心功能
# ============================================================
add_h1("四、核心功能")

add_h2("4.1 今日聚焦 & 三只青蛙")
add_bullet("跨项目汇总今天到期 / 逾期 / 进行中的任务，按项目分组")
add_bullet("顶部固定 3 件要事（MIT），复制清单时置顶输出")
add_bullet("一键复制今日清单，直接发团队群")

add_h2("4.2 任务四种视图")
add_image(imgs["views"], width=6.2, caption="图：列表 / 看板 / 日历 / 时间线")
add_table(
    ["视图", "快捷键", "用途"],
    [["☰ 列表", "1", "搜索 / 筛选，逐条管理"],
     ["▦ 看板", "2", "拖拽改状态，WIP 超限提醒"],
     ["📅 日历", "3", "月历看截止日分布，循环任务虚线占位"],
     ["⏱ 时间线", "4", "甘特图，日/周/月缩放，拖条改期，导出 PNG"]],
    col_widths=[1.5, 1.0, 4.0]
)

add_h2("4.3 快速添加语法")
add_body("一行文字搞定任务创建，自动解析日期、负责人、优先级和标记：")
add_image(imgs["quickadd"], width=5.5, caption="图：快速添加语法示例")
add_table(
    ["语法", "含义", "示例"],
    [["日期", "到期日", "今天 / 明天 / 后天 / 下周三 / +3天 / 2026-09-14"],
     ["@", "负责人", "@张三"],
     ["!", "优先级", "!P0（最高）/ !P1 / !P2"],
     ["#", "标记", "#风险 #每日 #每周 #每月"],
     ["📥", "先进收件箱", "输入框前点 📥 图标，暂不分项目"]],
    col_widths=[1.0, 1.5, 4.0]
)

add_h2("4.4 收件箱分拣")
add_bullet("快速添加可先进收件箱不分项目")
add_bullet("侧栏入口逐条分派到对应项目（Linear Triage 思路）")

add_h2("4.5 检查清单 & 进度上卷")
add_bullet("任务可挂子任务，清单完成率即任务进度条")
add_bullet("父任务进度自动上滚汇总（GitHub sub-issues 思路）")

add_h2("4.6 一键推迟 & 拖延徽章")
add_bullet("卡片上「⏭ 明天」一键推迟")
add_bullet("同一任务推迟 ≥ 3 次亮「⚠ 已推迟 N 次」逼出决策")

add_h2("4.7 循环任务预占日历")
add_bullet("每周 / 每月循环任务在日历画虚线占位块（Reclaim.ai 简化版）")

add_h2("4.8 Ctrl+K 命令面板")
add_bullet("跨项目搜索项目 / 任务，键盘直达")
add_bullet("新增 ➕ 新建任务 / 🌙 每日收尾问答 / 🌓 切换主题 / ⚙ 打开设置")

# ============================================================
#  五、项目管理
# ============================================================
add_h1("五、项目管理")

add_h2("5.1 多项目管理")
add_bullet("项目归档（已归档可折叠展开）、项目设置")
add_bullet("里程碑倒计时显示在项目工作区顶部")

add_h2("5.2 项目模板")
add_bullet("「另存为模板」保存当前项目任务集骨架")
add_bullet("新建项目时一键带入模板，快速启动同类项目")

add_h2("5.3 项目工具菜单（⋯）")
add_table(
    ["功能", "说明"],
    [["Excel 导入向导", "五步导入：选文件→选表→列映射+值映射→预览→导入报告"],
     ["导出 .xlsx 计划表", "按阶段分组，带样式，发甲方 / 领导"],
     ["导出 CSV", "通用表格格式"],
     ["导出 Markdown", "每任务一文件 + 项目主页，Obsidian 可接管"],
     ["项目设置", "企业微信/钉钉/飞书 Webhook、预算工时等"],
     ["另存为模板", "保存为项目模板"]],
    col_widths=[2.0, 4.5]
)

# ============================================================
#  六、汇报与沟通
# ============================================================
add_h1("六、汇报与沟通")

add_h2("6.1 风险登记册")
add_bullet("风险从布尔标记升级为 概率(1-5) × 影响(1-5) = 风险值")
add_bullet("≥ 15 标红置顶，强制进日报风险栏")
add_bullet("附应对措施 / 升级条件，跨项目登记册表格视图")

add_h2("6.2 每日收尾问答")
add_bullet("每天 17:30（可配）弹三问：今天干成什么 / 卡在哪 / 明天三件事")
add_bullet("次日日报自动引用（Sunsama Daily Shutdown 思路）")
add_bullet("新增「阻碍」一问，一键转风险任务")

add_h2("6.3 日报 / 周报一键发群")
add_bullet("项目设置里配企业微信 / 钉钉（支持加签）/ 飞书群机器人 Webhook")
add_bullet("报告一键推送到群，日报附工时、风险、最近决策")
add_callout("日报 / 周报自动汇总本周各项目投入工时，统计页有工时分布", color="E76F51", icon="📊")

# ============================================================
#  七、洞察与统计
# ============================================================
add_h1("七、洞察与统计")

add_h2("7.1 跨项目智能视图")
add_bullet("内置「风险登记册 / 风险且逾期 / 等待超 3 天」")
add_bullet("可保存自定义筛选（Vikunja saved filters 思路）")

add_h2("7.2 周期统计 Cycle Time")
add_bullet("记录开始 / 完成时间，统计平均处理时长与近 4 周逾期率趋势")

add_h2("7.3 需求池 RICE")
add_bullet("想法 / 需求按 价值 ÷ 工时 打分排序")
add_bullet("四象限决策，一键转为任务（airfocus / Productboard 思路）")

add_h2("7.4 蒙特卡洛交付预测")
add_bullet("50% / 85% / 95% 三档置信日期 + 直方图 + 一键复制话术")
add_bullet("小数据自动按周降级口径")

add_h2("7.5 其他统计")
add_bullet("吞吐量柱状图（中位数线）、SLE 分位阈值（85 分位，看板 WIP 超龄标红）")
add_bullet("工时预算告警（项目设预算小时数，80% / 100% 托盘通知 + 群推送）")

# ============================================================
#  八、自动化与 AI 集成
# ============================================================
add_h1("八、自动化与 AI 集成")

add_h2("8.1 固定规则自动化")
add_table(
    ["规则", "说明"],
    [["逾期自动置顶", "到期未完成的任务自动置顶到今日聚焦"],
     ["停滞进待清理", "停滞 14 天的任务进入待清理视图"],
     ["完成自动停表", "任务完成时自动停止计时器"],
     ["收件箱提醒", "每日提醒附带收件箱未分拣数量"]],
    col_widths=[2.0, 4.5]
)

add_h2("8.2 AI 提效三件套（可选）")
add_callout("AI 只产草稿，不直接改数据。确认后才入库。", color="E76F51", icon="⚠")
add_h3("供应商管理")
add_bullet("保存多家供应商档案，点卡片一键切换")
add_bullet("预设：智谱 GLM / DeepSeek / Kimi / 阿里云百炼 / SiliconFlow / OpenAI / OpenRouter / 本机 Ollama")
add_bullet("或任何 OpenAI 兼容服务；拉取模型列表点击选型、测试连接")
add_bullet("各档案独立保存 API Key（仅存本机 SQLite，本地进程直连供应商）")
add_h3("三大功能")
add_table(
    ["功能", "说明"],
    [["🤖 粘贴转任务", "会议记录 / 群聊 / 邮件整段粘入，自动拆成任务草稿（日期 / 负责人 / 优先级 / 风险）"],
     ["🤖 生成检查清单", "按任务标题拆出可验证步骤草稿"],
     ["日报 AI 润色", "一键润色日报，支持撤销"]],
    col_widths=[1.8, 4.7]
)

add_h2("8.3 每日定时提醒")
add_bullet("到点系统通知今日到期任务（默认 09:30，左下角 ⏰ 可改）")

add_h2("8.4 系统集成")
add_bullet("系统托盘常驻、关闭窗口最小化到托盘")
add_bullet("启动时通知今日到期任务、窗口状态记忆")
add_bullet("全局快速捕获：Alt+Shift+A 弹无边框小窗，默认进收件箱")
add_bullet("深链接 pm-todo://task/123 + 单实例：群消息点链接跳回任务")
add_bullet("任务栏进度条（专注会话进度 / 今日完成 vs 目标）")

# ============================================================
#  九、组织记忆
# ============================================================
add_h1("九、组织记忆")

add_h2("9.1 项目决策日志")
add_bullet("ADR 模板（架构决策记录），Ctrl+K 一键记录")
add_bullet("日报自动附最近决策")

add_h2("9.2 会议纪要 → 行动项")
add_bullet("会后笔记模板，行动项一键转任务（可溯源）")
add_bullet("一键进决策日志")

add_h2("9.3 干系人跟进")
add_bullet("联系人档案 + 跟进周期")
add_bullet("到期自动生成「跟进某人」进今日聚焦")

add_h2("9.4 每日笔记")
add_bullet("当日任务 + 青蛙 + 收尾 + 会议自动聚合")

add_h2("9.5 知识深度")
add_bullet("备注 Markdown 化（列表 / 粗体 / 引用，卡片与详情渲染）")
add_bullet("[[反链]] 索引：详情页显示「哪些任务/决策提到过它」")
add_bullet("项目 one-pager 档案：一页纸 + 自动嵌入任务表 / 风险册 / 预测 / 统计")
add_bullet("专注会话：全屏秒表 / 倒计时，结束自动记工时 + 一句收尾进每日笔记")

# ============================================================
#  十、数据安全
# ============================================================
add_h1("十、数据安全")

add_h2("10.1 回收站 / 软删除")
add_bullet("删除的任务 / 项目进回收站保留 30 天，可恢复")

add_h2("10.2 自动备份")
add_bullet("本机 SQLite 存储，启动自动备份（保留最近 30 份）")
add_bullet("支持 JSON 导入 / 导出、CSV 导出")

add_h2("10.3 数据操作入口")
add_table(
    ["操作", "入口位置"],
    [["立即备份", "⚙ 设置 → 💾 数据"],
     ["导出 JSON", "⚙ 设置 → 💾 数据"],
     ["导入 JSON", "⚙ 设置 → 💾 数据"],
     ["回收站", "侧栏页脚 🗑 / ⚙ 设置 → 💾 数据"],
     ["Excel 导入", "项目工作区顶栏 ⋯ → Excel 导入向导"],
     ["导出 .xlsx / CSV / Markdown", "项目工作区顶栏 ⋯ → 对应选项"]],
    col_widths=[2.5, 4.0]
)

# ============================================================
#  十一、快捷键速查
# ============================================================
add_h1("十一、快捷键速查")
add_table(
    ["快捷键", "功能"],
    [["N", "新建任务"],
     ["1", "列表视图"],
     ["2", "看板视图"],
     ["3", "日历视图"],
     ["4", "时间线（甘特）视图"],
     ["5", "统计页签"],
     ["6", "决策页签"],
     ["7", "档案页签"],
     ["/", "搜索"],
     ["Ctrl + K", "命令面板（搜项目/任务 + 直达操作）"],
     ["Alt + Shift + A", "全局快速捕获（弹小窗，默认进收件箱）"]],
    col_widths=[2.0, 4.5]
)
add_callout("快捷键 1-4 切换任务四种视图，5-7 直达统计/决策/档案，与界面顺序完全一致。", color="2A9D8F")

# ============================================================
#  十二、从源码构建
# ============================================================
add_h1("十二、从源码构建")
add_body("依赖：Rust（MSVC 工具链）、WebView2（Win10/11 通常已内置）。")

add_h3("构建步骤")
add_body("1. 安装 Rust（也可使用仓库 tools/install-rust.cmd 便携方案）", color="555555")
add_body("   rustup default stable", color="2A9D8F")
add_body("2. 安装 Tauri CLI", color="555555")
add_body("   cargo install tauri-cli", color="2A9D8F")
add_body("3. 开发调试", color="555555")
add_body("   cargo tauri dev", color="2A9D8F")
add_body("4. 打包 Windows 安装包（NSIS）", color="555555")
add_body("   cargo tauri build", color="2A9D8F")
add_body("产物位于 src-tauri/target/release/bundle/nsis/", size=9.5, color="8A99A8")

add_h3("前端开发预览")
add_body("浏览器预览（mock 后端，数据存 localStorage）：")
add_body("   node tools/serve-ui.js  →  http://127.0.0.1:8123/", color="2A9D8F")
add_callout("ES Modules 不能通过 file:// 直开，需用 serve-ui.js 启动本地服务。", color="E76F51", icon="⚠")

add_h3("发布新版本")
add_body("推送 v 开头的标签即自动在 GitHub Actions 编译发布 Release：")
add_body("   git tag v2.0.0", color="2A9D8F")
add_body("   git push origin v2.0.0", color="2A9D8F")
add_body("也可在仓库 Actions 页面手动触发，版本号自动取自 tauri.conf.json。", size=9.5, color="8A99A8")

# ── 页脚 ──
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("PM 待办助手  ·  MIT License  ·  github.com/wangjie-git/pm_tool")
_set_run(r, size=9, color="8A99A8")

# ============================================================
out = os.path.join(WORK, "PM待办助手-使用手册.docx")
doc.save(out)
print(f"\n手册已生成: {out}")
print(f"文件大小: {round(os.path.getsize(out)/1024,1)} KB")
