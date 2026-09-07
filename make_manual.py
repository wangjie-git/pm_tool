# -*- coding: utf-8 -*-
"""生成 OpenSquilla 工具使用手册（图文并茂，Word 格式）。"""
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
import numpy as np

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"D:\ai\pm-todo"
ASSET = os.path.join(BASE, "assets")
os.makedirs(ASSET, exist_ok=True)

# 中文字体
plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei"]
plt.rcParams["axes.unicode_minus"] = False

# 主题配色（皮皮虾主题：青+橙+珊瑚）
C_TEAL = "#0E7C7B"
C_ORANGE = "#F28C28"
C_CORAL = "#E8594F"
C_DEEP = "#173B5C"
C_GRID = "#E4ECF2"
PALETTE = ["#0E7C7B", "#F28C28", "#E8594F", "#2B8FB3", "#7A6FBE",
           "#3BA55D", "#D28B2F", "#C94F7C", "#5C8D5C", "#4C6FA8"]


def emo():
    return ["💬", "📁", "⚙️", "🌐", "📄", "🎙️", "🧠", "🧩", "⏰", "🖼️"]


CATEGORIES = [
    ("会话与消息", 9),
    ("文件与工作区", 8),
    ("代码执行", 4),
    ("网络检索", 4),
    ("文档与交付", 5),
    ("音频与语音", 11),
    ("记忆", 2),
    ("技能", 8),
    ("定时任务", 1),
    ("其他工具", 4),
]


# ---------- 图 1：分类统计柱状图 ----------
def make_bar():
    cats = [c[0] for c in CATEGORIES]
    vals = [c[1] for c in CATEGORIES]
    fig, ax = plt.subplots(figsize=(8.4, 4.4), dpi=160)
    y = np.arange(len(cats))[::-1]
    bars = ax.barh(y, vals, color=PALETTE, height=0.62, zorder=3)
    for yi, v in zip(y, vals):
        ax.text(v + 0.18, yi, str(v), va="center", fontsize=12,
                fontweight="bold", color=C_DEEP)
    ax.set_yticks(y)
    ax.set_yticklabels(cats, fontsize=12)
    ax.set_xlim(0, max(vals) + 1.6)
    ax.set_xlabel("工具数量（个）", fontsize=11, color=C_DEEP)
    ax.set_title("OpenSquilla 工具分类一览", fontsize=15, fontweight="bold",
                 color=C_DEEP, pad=14)
    ax.xaxis.set_visible(False)
    for s in ["top", "right", "bottom", "left"]:
        ax.spines[s].set_visible(False)
    ax.grid(axis="x", color=C_GRID, zorder=0)
    ax.tick_params(axis="y", length=0)
    fig.tight_layout()
    fig.savefig(os.path.join(ASSET, "cat_bar.png"), transparent=True)
    plt.close(fig)


# ---------- 图 2：占比环形图 ----------
def make_donut():
    vals = np.array([c[1] for c in CATEGORIES], dtype=float)
    labels = [c[0] for c in CATEGORIES]
    fig, ax = plt.subplots(figsize=(5.6, 4.4), dpi=160)
    wedges, _ = ax.pie(vals, colors=PALETTE, startangle=90,
                       counterclock=False, wedgeprops=dict(width=0.42,
                       edgecolor="white", linewidth=2))
    total = int(vals.sum())
    ax.text(0, 0.06, f"{total}", ha="center", va="center", fontsize=26,
            fontweight="bold", color=C_DEEP)
    ax.text(0, -0.24, "工具总数", ha="center", va="center", fontsize=10,
            color="#6b7c8a")
    ax.legend(wedges, [f"{l}  {int(v)}" for l, v in zip(labels, vals)],
              loc="center left", bbox_to_anchor=(1.0, 0.5), fontsize=9,
              frameon=False)
    ax.set_title("分类占比", fontsize=14, fontweight="bold", color=C_DEEP)
    fig.tight_layout()
    fig.savefig(os.path.join(ASSET, "cat_donut.png"), transparent=True)
    plt.close(fig)


# ---------- 图 3：工作流程图 ----------
def make_workflow():
    fig, ax = plt.subplots(figsize=(9.2, 3.6), dpi=160)
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 34)
    ax.axis("off")

    stages = [
        ("01 明确需求", "理解任务 · 读取上下文", C_TEAL),
        ("02 检索与记忆", "web_search · memory_search", C_ORANGE),
        ("03 执行与计算", "exec_command · execute_code", C_CORAL),
        ("04 生成与交付", "create_* · publish_artifact", "#7A6FBE"),
        ("05 收尾", "验证结果 · 汇报", "#3BA55D"),
    ]
    n = len(stages)
    box_w, box_h = 16.5, 15.0
    gap = (100 - box_w * n) / (n + 1)
    ys = 20
    xs = []
    for i, (title, sub, color) in enumerate(stages):
        x = gap + i * (box_w + gap)
        xs.append(x)
        ax.add_patch(FancyBboxPatch((x, ys - box_h / 2), box_w, box_h,
                     boxstyle="round,pad=0.6,rounding_size=1.4",
                     fc=color, ec="none", zorder=3))
        ax.text(x + box_w / 2, ys + 2.6, title, ha="center", va="center",
                fontsize=13, fontweight="bold", color="white", zorder=4)
        ax.text(x + box_w / 2, ys - 2.4, sub, ha="center", va="center",
                fontsize=8.6, color="white", zorder=4)
        if i < n - 1:
            ax.add_patch(FancyArrowPatch(
                (x + box_w + 0.5, ys), (x + box_w + gap - 0.5, ys),
                arrowstyle="-|>", mutation_scale=18, lw=2.2,
                color=C_DEEP, zorder=3))

    ax.text(50, 31.5, "一次典型任务的执行流水线", ha="center", va="center",
            fontsize=14, fontweight="bold", color=C_DEEP)
    fig.tight_layout()
    fig.savefig(os.path.join(ASSET, "workflow.png"), transparent=True)
    plt.close(fig)


# ---------- 图 4：封面横幅 ----------
def make_banner():
    fig, ax = plt.subplots(figsize=(8.6, 2.6), dpi=160)
    x = np.linspace(0, 1, 256)
    grad = np.vstack([x, x])
    ax.imshow(grad, extent=[0, 1, 0, 1], aspect="auto",
              cmap="viridis", zorder=0)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.text(0.5, 0.66, "OpenSquilla 工具使用手册", ha="center",
            va="center", fontsize=26, fontweight="bold", color="white")
    ax.text(0.5, 0.32, "Tool & Skill · 快速上手指南", ha="center",
            va="center", fontsize=13, color="white", alpha=0.92)
    fig.tight_layout(pad=0)
    fig.savefig(os.path.join(ASSET, "banner.png"), transparent=False,
                facecolor="#173B5C")
    plt.close(fig)


for fn in [make_bar, make_donut, make_workflow, make_banner]:
    fn()
print("图片生成完成:", os.listdir(ASSET))


# ================= 构建 Word =================
TOOLS = {
    "会话与消息": [
        ("sessions_list", "列出活跃会话（可按状态/代理过滤）"),
        ("sessions_history", "读取某会话的历史对话"),
        ("sessions_spawn", "派生隔离的子代理会话"),
        ("sessions_send", "向另一会话发送消息（代理间通信）"),
        ("sessions_yield", "挂起当前回合，等待子代理完成"),
        ("session_search", "全文检索历史会话记录"),
        ("session_status", "查看当前会话用量/成本/模型"),
        ("agents_list", "列出可用的代理配置"),
        ("message", "经渠道适配器向外部用户/群聊发送消息"),
    ],
    "文件与工作区": [
        ("list_dir", "列出目录内容（类型/大小）"),
        ("glob_search", "按通配符模式查找文件"),
        ("grep_search", "按正则表达式检索文件内容"),
        ("read_file", "读取 UTF-8 文本文件"),
        ("write_file", "写入完整文件内容"),
        ("edit_file", "精确文本替换编辑文件"),
        ("apply_patch", "结构化补丁（增/删/改多行）"),
        ("read_spreadsheet", "读取 CSV/TSV/Excel 表格数据"),
    ],
    "代码执行": [
        ("exec_command", "前台执行 shell 命令"),
        ("execute_code", "在隔离子进程执行 Python"),
        ("background_process", "后台运行 shell 命令"),
        ("process", "管理后台进程（等待/日志/终止）"),
    ],
    "网络检索": [
        ("web_search", "带来源引用的联网搜索"),
        ("web_discover", "轻量链接发现（标题+摘要）"),
        ("web_fetch", "抓取网页并提取可读正文"),
        ("http_request", "发起 HTTP 请求（可保存响应）"),
    ],
    "文档与交付": [
        ("create_csv", "生成 CSV 文件并交付"),
        ("create_xlsx", "生成 Excel 工作簿并交付"),
        ("create_pdf_report", "生成 PDF 报告并交付"),
        ("pdf", "提取 PDF 文本内容"),
        ("publish_artifact", "注册工作区文件为交付物"),
    ],
    "音频与语音": [
        ("tts", "文本转语音合成"),
        ("audio_config", "配置语音服务商（热生效）"),
        ("audio_provider_capabilities", "查询语音服务能力"),
        ("voice_search", "搜索共享音色"),
        ("voice_clone", "从音频样本克隆音色"),
        ("voice_convert", "将源音频转换到目标音色"),
        ("dubbing_generate", "提交配音任务"),
        ("dubbing_status", "查询配音任务状态"),
        ("dubbing_download", "下载配音音频"),
        ("music_generate", "生成器乐"),
        ("song_generate", "生成带人声的歌曲"),
    ],
    "记忆": [
        ("memory_search", "检索长期记忆/历史决策"),
        ("memory_get", "读取记忆源文件内容"),
    ],
    "技能": [
        ("skill_list", "列出可用技能"),
        ("skill_view", "查看技能内容"),
        ("skill_create", "创建本地技能"),
        ("skill_edit", "编辑技能"),
        ("skill_delete", "删除技能"),
        ("skill_install_community", "安装社区技能"),
        ("skill_search_community", "搜索社区技能市场"),
        ("install_skill_deps", "安装技能依赖"),
    ],
    "定时任务": [
        ("cron", "创建/列出/删除/触发定时任务"),
    ],
    "其他工具": [
        ("image", "用视觉模型分析图片"),
        ("router_control", "切换路由/模型"),
        ("gateway", "读取网关配置"),
        ("retrieve_tool_result", "取回被省略的原始输出"),
    ],
}

CAT_EMOJI = {
    "会话与消息": "💬", "文件与工作区": "📁", "代码执行": "⚙️",
    "网络检索": "🌐", "文档与交付": "📄", "音频与语音": "🎙️",
    "记忆": "🧠", "技能": "🧩", "定时任务": "⏰", "其他工具": "🖼️",
}

CAT_TIPS = {
    "会话与消息": "多代理协作时，用 sessions_spawn 派生子代理、sessions_send 互通消息，最后 sessions_yield 收回结果。",
    "文件与工作区": "读取后再修改：改文件前先用 read_file 建立编辑上下文，再用 edit_file / apply_patch 做精确替换。",
    "代码执行": "长任务用 background_process 后台跑，配合 process 的 wait 阻塞等待，避免反复轮询。",
    "网络检索": "需要“当前信息+出处”优先 web_search；只想要链接列表用 web_discover；抓指定页面用 web_fetch。",
    "文档与交付": "文件型成果（表格/PDF/报告）生成后务必 publish_artifact 注册，交界面会自动提供下载。",
    "音频与语音": "配音/克隆需显式同意（consent）；audio_config 配置后即时生效，无需重启网关。",
    "记忆": "身份/偏好放 USER.md；长期事实放 MEMORY.md；每日随笔放 memory/日期.md，用 memory_search 检索。",
    "技能": "按需启用：先 skill_view 查看玩法，再决定是否 skill_install_community 安装。",
    "定时任务": "提醒类用 job_kind=reminder 直接投递，无需模型；后台代理任务才用 agent_turn。",
    "其他工具": "image 用于读图分析；router_control 仅在用户明确要求切换路由时使用。",
}

WORKFLOW_EXAMPLES = [
    ("查资料 → 出报告", "web_search 检索 → 汇总要点 → create_pdf_report 生成 → publish_artifact 交付"),
    ("处理数据 → 出表格", "read_spreadsheet 读表 → execute_code 计算 → create_xlsx 导出 → 发布"),
    ("写代码 → 跑测试", "write_file 写脚本 → exec_command 运行 → grep_search 排查 → 修正并交付"),
    ("做配音", "dubbing_generate 提交 → dubbing_status 轮询 → dubbing_download 下载音频"),
]


def set_cn_font(run, name="微软雅黑", size=None, bold=None, color=None, italic=None):
    run.font.name = "Calibri"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    if italic is not None:
        run.font.italic = italic


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.space_before = Pt(0)
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        set_cn_font(run, size=16, bold=True, color="0E7C7B")
    elif level == 2:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        set_cn_font(run, size=13, bold=True, color="173B5C")
    else:
        run = p.add_run(text)
        set_cn_font(run, size=11, bold=True, color="F28C28")
    # 底部边框
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6" if level == 1 else "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), (C_TEAL if level == 1 else "C9D7E0").lstrip("#"))
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


def add_para(doc, text, size=10.5, color="2A3B4C", bold=False, space_after=4,
             italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_cn_font(run, size=size, bold=bold, color=color)
    if italic:
        run.font.italic = True
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_cn_font(run, size=10.5, color="2A3B4C")
    return p


def shade_cell(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor.lstrip("#"))
    tcPr.append(shd)


def add_tool_table(doc, tools):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, txt in enumerate(["工具", "用途"]):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(txt)
        set_cn_font(run, size=10, bold=True, color="FFFFFF")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade_cell(hdr[i], C_TEAL)
    for name, desc in tools:
        row = t.add_row().cells
        row[0].text = ""
        p0 = row[0].paragraphs[0]
        r0 = p0.add_run(name)
        set_cn_font(r0, size=9.5, bold=True, color="0E7C7B")
        # 等宽效果
        r0.font.name = "Consolas"
        row[1].text = ""
        p1 = row[1].paragraphs[0]
        r1 = p1.add_run(desc)
        set_cn_font(r1, size=9.5, color="2A3B4C")
    # 列宽
    for row in t.rows:
        row.cells[0].width = Inches(2.4)
        row.cells[1].width = Inches(4.3)
    return t


doc = Document()
# 页面边距
for sec in doc.sections:
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)

# Normal 样式默认字体
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

# ===== 封面 =====
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(20)
run = p.add_run("🦐")
set_cn_font(run, size=30)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("OpenSquilla 工具使用手册")
set_cn_font(r, size=26, bold=True, color=C_DEEP)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Tool & Skill · 图文并茂 · 言简意赅")
set_cn_font(r, size=12, color="6b7c8a")

doc.add_picture(os.path.join(ASSET, "banner.png"), width=Inches(6.6))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("版本 1.0 · 生成日期 2026-09-06")
set_cn_font(r, size=10, color="8a9aa8")

doc.add_page_break()

# ===== 1 简介 =====
add_heading(doc, "一、简介", 1)
add_para(doc, "OpenSquilla 是一款面向智能体的运行环境，内置十余类、共 50+ 个工具，"
              "覆盖会话管理、文件操作、代码执行、网络检索、文档交付、音视频生成、记忆与定时任务等。"
              "本手册按“先总览、后分类、再实战”的顺序编排，助你快速上手。", size=10.5)
add_para(doc, "阅读建议：", size=10.5, bold=True, color=C_DEEP, space_after=2)
add_bullet(doc, "只做一件小事 → 直接查对应分类的工具表。")
add_bullet(doc, "想串成完整任务 → 看“常用工作流”。")
add_bullet(doc, "需要长期复用 → 用“技能 / 记忆”沉淀。")

# ===== 2 工具全景 =====
add_heading(doc, "二、工具全景", 1)
add_para(doc, "全部工具按功能分为 10 大类，音频/语音类最丰富（11 个），其次是会话与消息、文件、技能等。", size=10.5)

pic1 = doc.add_paragraph()
pic1.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_picture(os.path.join(ASSET, "cat_bar.png"), width=Inches(6.3))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

pic2 = doc.add_paragraph()
pic2.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_picture(os.path.join(ASSET, "cat_donut.png"), width=Inches(4.4))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

# ===== 3 分类详解 =====
add_heading(doc, "三、分类详解", 1)
for i, (cat, _) in enumerate(CATEGORIES):
    add_heading(doc, f"{CAT_EMOJI[cat]}  {cat}", 2)
    add_tool_table(doc, TOOLS[cat])
    tip = doc.add_paragraph()
    tip.paragraph_format.space_before = Pt(4)
    tip.paragraph_format.space_after = Pt(6)
    r = tip.add_run("💡 " + CAT_TIPS[cat])
    set_cn_font(r, size=9.5, italic=True, color="B06A1E")

# ===== 4 常用工作流 =====
add_heading(doc, "四、常用工作流", 1)
doc.add_picture(os.path.join(ASSET, "workflow.png"), width=Inches(6.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_para(doc, "典型场景示例：", size=10.5, bold=True, color=C_DEEP, space_after=2)
for title, flow in WORKFLOW_EXAMPLES:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title + "：")
    set_cn_font(r, size=10, bold=True, color=C_DEEP)
    r2 = p.add_run(flow)
    set_cn_font(r2, size=10, color="2A3B4C")

# ===== 5 注意事项 =====
add_heading(doc, "五、使用建议与注意事项", 1)
for tip in [
    "安全第一：不绕过/弱化任何安全措施；涉及音色克隆、配音需取得说话人同意。",
    "先读后改：编辑现有文件前先 read_file，避免丢失原有格式。",
    "数据可溯：表格统计、TopN、求和等用 execute_code / read_spreadsheet 求真实值，不要凭记忆估算。",
    "交付闭环：文件型成果生成后调用 publish_artifact 注册，交界面会提供下载入口。",
    "记忆分层：身份偏好 → USER.md；长期事实 → MEMORY.md；每日随笔 → memory/日期.md。",
    "术语克制：对外交付/提交说明只描述产品改动与理由，不提及内部工具链名称。",
]:
    add_bullet(doc, tip)

footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.space_before = Pt(12)
r = footer.add_run("— 完 · OpenSquilla 🦐 —")
set_cn_font(r, size=10, color="8a9aa8")

OUT = os.path.join(BASE, "OpenSquilla工具使用手册.docx")
doc.save(OUT)
print("手册已生成:", OUT)
